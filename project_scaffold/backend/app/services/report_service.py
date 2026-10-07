import io
import logging
import json
from datetime import datetime
from typing import Dict, Any, List, Optional
from sqlalchemy.orm import Session

from app.models import (
    User, SkinAssessment, SkinHealthScore, SkincareRoutine, Recommendation,
    SkinProfile, LifestyleProfile, RoutineAdherenceRecord, ProgressRecord,
    ProductRecommendation, Product, HydrationRecord, SleepRecord, EnvironmentalExposure
)
from app.schemas.skin_intelligence import (
    SkinIntelligenceReport, ScoreComponentDetail, RoutineStep, RecommendationResponse
)
from app.services.skin_assessment_service import analyze_skin_profile
from app.services.scoring_service import calculate_skin_health_score
from app.services.routine_generation_service import generate_personalized_routine
from app.services.recommendation_service import generate_evidence_recommendations
from app.services.safety_service import evaluate_ingredient_safety
from app.services.evidence_service import get_all_ingredient_evidence

# ReportLab Imports
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# Openpyxl Imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

def generate_skin_intelligence_report(db: Session, user: User) -> SkinIntelligenceReport:
    """Generate the full 18-point clinical skin intelligence report."""
    latest_assessment = db.query(SkinAssessment).filter(SkinAssessment.user_id == user.id).order_by(SkinAssessment.created_at.desc()).first()
    if not latest_assessment:
        latest_assessment = analyze_skin_profile(db, user)

    latest_score = db.query(SkinHealthScore).filter(SkinHealthScore.user_id == user.id).order_by(SkinHealthScore.created_at.desc()).first()
    if not latest_score or latest_score.assessment_id != latest_assessment.id:
        latest_score = calculate_skin_health_score(db, user, latest_assessment)

    latest_routine = db.query(SkincareRoutine).filter(SkincareRoutine.user_id == user.id).order_by(SkincareRoutine.created_at.desc()).first()
    if not latest_routine or latest_routine.assessment_id != latest_assessment.id:
        latest_routine = generate_personalized_routine(db, user, latest_assessment)

    recs = db.query(Recommendation).filter(Recommendation.user_id == user.id, Recommendation.assessment_id == latest_assessment.id).all()
    if not recs:
        recs = generate_evidence_recommendations(db, user, latest_assessment)

    skin_prof: Optional[SkinProfile] = user.skin_profile
    life_prof: Optional[LifestyleProfile] = user.lifestyle_profile

    sec_concerns = json.loads(latest_assessment.secondary_concerns) if latest_assessment.secondary_concerns else []
    risk_factors = json.loads(latest_assessment.risk_factors) if latest_assessment.risk_factors else []
    positive_factors = json.loads(latest_assessment.supporting_factors) if latest_assessment.supporting_factors else []

    raw_breakdown = json.loads(latest_score.score_breakdown) if latest_score.score_breakdown else []
    breakdown_components = [ScoreComponentDetail(**b) for b in raw_breakdown]

    morning_steps = [RoutineStep(**s) for s in json.loads(latest_routine.morning_routine)]
    evening_steps = [RoutineStep(**s) for s in json.loads(latest_routine.evening_routine)]
    weekly_steps = [RoutineStep(**s) for s in json.loads(latest_routine.weekly_routine)]

    ing_recs = [RecommendationResponse.model_validate(r) for r in recs if r.recommendation_type == "INGREDIENT"]
    cat_recs = [RecommendationResponse.model_validate(r) for r in recs if r.recommendation_type == "PRODUCT_CATEGORY"]

    why_recs = [
        {"recommendation": r.ingredient_or_category, "reason": r.reason, "trigger": r.user_factor_trigger or "Skin Assessment Target"}
        for r in recs
    ]

    evidence_refs = [
        {"recommendation": r.ingredient_or_category, "evidence": r.evidence_reference or "Clinical Dermatology Practice Guidelines"}
        for r in recs
    ]

    all_evidence = get_all_ingredient_evidence()
    allergies = skin_prof.allergies if skin_prof else ""
    sensitivities = skin_prof.sensitivities if skin_prof else ""
    skin_type = skin_prof.skin_type.value if (skin_prof and skin_prof.skin_type) else "NORMAL"
    safety_eval = evaluate_ingredient_safety(all_evidence, allergies, sensitivities, skin_type)

    safety_notes = list(safety_eval.safety_notes)
    for exc in safety_eval.excluded_ingredients:
        safety_notes.append(f"Excluded '{exc['ingredient_name']}': {exc['reason']}")

    if not safety_notes:
        safety_notes.append("No active safety exclusions or allergy conflicts identified.")

    next_steps = [
        "Follow your personalized Morning, Evening, and Weekly routines consistently for 14 days.",
        "Track daily hydration and sleep restoration metrics to optimize your 5-part Skin Health Score.",
        "Log any newly observed sensitivities or changes in skin reaction in your profile.",
        "Schedule a follow-up assessment in 30 days to measure dermal progress."
    ]

    return SkinIntelligenceReport(
        user_summary={
            "id": user.id,
            "name": user.name,
            "email": user.email,
            "age": user.age or "Not specified",
            "country": user.country or "Not specified",
            "role": user.role.value
        },
        skin_profile={
            "skin_type": skin_prof.skin_type.value if (skin_prof and skin_prof.skin_type) else "Not configured",
            "reported_concerns": skin_prof.skin_concerns if skin_prof else "None",
            "allergies": skin_prof.allergies if (skin_prof and skin_prof.allergies) else "None recorded",
            "sensitivities": skin_prof.sensitivities if (skin_prof and skin_prof.sensitivities) else "None recorded"
        },
        overall_skin_health_score=latest_score.total_score,
        score_breakdown=breakdown_components,
        primary_concern=latest_assessment.primary_concern,
        secondary_concerns=sec_concerns,
        risk_factors=risk_factors,
        positive_factors=positive_factors,
        morning_routine=morning_steps,
        evening_routine=evening_steps,
        weekly_routine=weekly_steps,
        ingredient_recommendations=ing_recs,
        product_category_recommendations=cat_recs,
        why_these_recommendations=why_recs,
        safety_notes=safety_notes,
        evidence_references=evidence_refs,
        assessment_confidence=round(latest_assessment.confidence_score * 100, 1),
        data_completeness=round(latest_assessment.data_completeness * 100, 1),
        confidence_explanation=f"Based on {round(latest_assessment.data_completeness * 100)}% data completeness across profile and biometric tracking logs.",
        recommended_next_steps=next_steps,
        generated_at=datetime.utcnow()
    )


def generate_assessment_report(db: Session, user: User) -> dict:
    """Generate focused Skin Assessment report."""
    assessment = db.query(SkinAssessment).filter(SkinAssessment.user_id == user.id).order_by(SkinAssessment.created_at.desc()).first()
    if not assessment:
        assessment = analyze_skin_profile(db, user)
    return {
        "report_type": "SKIN_ASSESSMENT",
        "user_id": user.id,
        "name": user.name,
        "primary_concern": assessment.primary_concern,
        "secondary_concerns": json.loads(assessment.secondary_concerns or "[]"),
        "risk_factors": json.loads(assessment.risk_factors or "[]"),
        "supporting_factors": json.loads(assessment.supporting_factors or "[]"),
        "confidence_score": round(assessment.confidence_score * 100, 1),
        "data_completeness": round(assessment.data_completeness * 100, 1),
        "created_at": assessment.created_at.isoformat()
    }


def generate_aurascore_report(db: Session, user: User) -> dict:
    """Generate 5-factor AuraScore report."""
    score = db.query(SkinHealthScore).filter(SkinHealthScore.user_id == user.id).order_by(SkinHealthScore.created_at.desc()).first()
    if not score:
        assessment = db.query(SkinAssessment).filter(SkinAssessment.user_id == user.id).order_by(SkinAssessment.created_at.desc()).first()
        if not assessment:
            assessment = analyze_skin_profile(db, user)
        score = calculate_skin_health_score(db, user, assessment)

    breakdown = json.loads(score.score_breakdown or "[]")
    impact_factors = json.loads(score.top_impact_factors or "[]")

    return {
        "report_type": "AURASCORE_5_FACTOR",
        "user_id": user.id,
        "name": user.name,
        "total_score": score.total_score,
        "skin_condition_score": score.skin_condition_score,
        "lifestyle_score": score.lifestyle_score,
        "sleep_score": score.sleep_score,
        "routine_consistency_score": score.routine_consistency_score,
        "hydration_score": score.hydration_score,
        "breakdown": breakdown,
        "top_impact_factors": impact_factors,
        "calculated_at": score.created_at.isoformat()
    }


def generate_routine_report(db: Session, user: User) -> dict:
    """Generate Personalized Routine report."""
    routine = db.query(SkincareRoutine).filter(SkincareRoutine.user_id == user.id).order_by(SkincareRoutine.created_at.desc()).first()
    if not routine:
        assessment = db.query(SkinAssessment).filter(SkinAssessment.user_id == user.id).order_by(SkinAssessment.created_at.desc()).first()
        if not assessment:
            assessment = analyze_skin_profile(db, user)
        routine = generate_personalized_routine(db, user, assessment)

    return {
        "report_type": "PERSONALIZED_ROUTINE",
        "user_id": user.id,
        "name": user.name,
        "morning_routine": json.loads(routine.morning_routine or "[]"),
        "evening_routine": json.loads(routine.evening_routine or "[]"),
        "weekly_routine": json.loads(routine.weekly_routine or "[]"),
        "seasonal_notes": routine.seasonal_notes,
        "safety_notes": routine.safety_notes,
        "created_at": routine.created_at.isoformat()
    }


def generate_product_recommendation_report(db: Session, user: User) -> dict:
    """Generate Product and Ingredient recommendations report."""
    recs = db.query(ProductRecommendation).filter(ProductRecommendation.user_id == user.id).all()
    products_data = []
    for pr in recs:
        p = db.query(Product).filter(Product.id == pr.product_id).first()
        if p:
            products_data.append({
                "product_name": p.name,
                "brand": p.brand,
                "category": p.category,
                "price": p.price,
                "suitability_score": pr.suitability_score,
                "match_reasons": json.loads(pr.reasons or "[]"),
                "is_current_routine": pr.is_current_routine
            })
    return {
        "report_type": "PRODUCT_RECOMMENDATIONS",
        "user_id": user.id,
        "total_recommendations": len(products_data),
        "products": products_data
    }


def generate_ingredient_safety_report(db: Session, user: User) -> dict:
    """Generate ingredient safety and contraindications analysis report."""
    skin_prof: Optional[SkinProfile] = user.skin_profile
    all_evidence = get_all_ingredient_evidence()
    allergies = skin_prof.allergies if skin_prof else ""
    sensitivities = skin_prof.sensitivities if skin_prof else ""
    skin_type = skin_prof.skin_type.value if (skin_prof and skin_prof.skin_type) else "NORMAL"
    safety_eval = evaluate_ingredient_safety(all_evidence, allergies, sensitivities, skin_type)

    return {
        "report_type": "INGREDIENT_SAFETY",
        "user_id": user.id,
        "skin_type": skin_type,
        "recorded_allergies": allergies,
        "recorded_sensitivities": sensitivities,
        "excluded_ingredients": safety_eval.excluded_ingredients,
        "safety_notes": list(safety_eval.safety_notes)
    }


def generate_progress_report(db: Session, user: User) -> dict:
    """Generate longitudinal progress report."""
    records = db.query(ProgressRecord).filter(ProgressRecord.user_id == user.id).order_by(ProgressRecord.recorded_at.asc()).all()
    adherences = db.query(RoutineAdherenceRecord).filter(RoutineAdherenceRecord.user_id == user.id).order_by(RoutineAdherenceRecord.date.asc()).all()

    avg_adherence = round((sum(a.adherence_rate for a in adherences) / len(adherences)) * 100, 1) if adherences else 0.0

    return {
        "report_type": "PROGRESS_LONGITUDINAL",
        "user_id": user.id,
        "total_progress_snapshots": len(records),
        "average_adherence_rate_pct": avg_adherence,
        "snapshots": [
            {
                "id": r.id,
                "total_score": r.total_score,
                "barrier_status": r.barrier_status,
                "recorded_at": r.recorded_at.isoformat()
            } for r in records
        ],
        "recent_adherence_records": [
            {
                "date": a.date,
                "morning_completed": a.morning_completed,
                "evening_completed": a.evening_completed,
                "adherence_rate": a.adherence_rate
            } for a in adherences[-14:]
        ]
    }


def generate_before_after_report(db: Session, user: User) -> dict:
    """Generate Before/After clinical comparison report."""
    records = db.query(ProgressRecord).filter(ProgressRecord.user_id == user.id).order_by(ProgressRecord.recorded_at.asc()).all()
    if not records:
        return {"report_type": "BEFORE_AFTER_COMPARISON", "status": "insufficient_data", "message": "At least one progress snapshot is required."}

    initial = records[0]
    latest = records[-1]

    score_delta = latest.total_score - initial.total_score

    return {
        "report_type": "BEFORE_AFTER_COMPARISON",
        "user_id": user.id,
        "initial_snapshot": {
            "date": initial.recorded_at.strftime("%Y-%m-%d"),
            "score": initial.total_score,
            "barrier_status": initial.barrier_status
        },
        "latest_snapshot": {
            "date": latest.recorded_at.strftime("%Y-%m-%d"),
            "score": latest.total_score,
            "barrier_status": latest.barrier_status
        },
        "overall_score_improvement": score_delta,
        "outcome": "IMPROVED" if score_delta > 0 else ("MAINTAINED" if score_delta == 0 else "ATTENTION_NEEDED")
    }


# =========================================================================
# PDF REPORT EXPORT GENERATION (ReportLab)
# =========================================================================

def generate_report_pdf(db: Session, user: User, report_type: str = "comprehensive") -> bytes:
    """
    Generates a beautifully styled, branded clinical PDF report using ReportLab.
    """
    report_data = generate_skin_intelligence_report(db, user)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom Color Palette
    PRIMARY = colors.HexColor("#0d9488")      # Teal 600
    PRIMARY_DARK = colors.HexColor("#0f766e") # Teal 700
    TEXT_DARK = colors.HexColor("#0f172a")    # Slate 900
    TEXT_MUTED = colors.HexColor("#475569")   # Slate 600
    BG_CARD = colors.HexColor("#f8fafc")      # Slate 50
    BORDER = colors.HexColor("#e2e8f0")       # Slate 200
    ACCENT_GOLD = colors.HexColor("#d97706")  # Amber 600
    ACCENT_GREEN = colors.HexColor("#16a34a") # Green 600

    # Custom Typography Styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=PRIMARY_DARK
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=TEXT_MUTED
    )
    heading_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=PRIMARY_DARK,
        spaceBefore=10,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'BodyText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=TEXT_DARK
    )
    bold_style = ParagraphStyle(
        'BoldText',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=12,
        textColor=TEXT_DARK
    )
    disclaimer_style = ParagraphStyle(
        'Disclaimer',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=7,
        leading=9.5,
        textColor=TEXT_MUTED
    )

    story = []

    # 1. Header Banner Table
    header_data = [
        [
            Paragraph("<b>AURASKIN</b> <font size=9 color='#0d9488'>AI SKIN INTELLIGENCE</font>", title_style),
            Paragraph(f"<b>CONFIDENTIAL CLINICAL REPORT</b><br/><font size=8 color='#64748b'>Generated: {report_data.generated_at.strftime('%b %d, %Y • %H:%M UTC')}</font>", ParagraphStyle('RAlign', parent=subtitle_style, alignment=2))
        ]
    ]
    header_table = Table(header_data, colWidths=[3.8 * inch, 3.8 * inch])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(header_table)
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY, spaceBefore=4, spaceAfter=8))

    # 2. Patient & Assessment Summary Card
    client_info = [
        [
            Paragraph("<b>Client Name:</b>", bold_style),
            Paragraph(str(report_data.user_summary.get("name", "N/A")), body_style),
            Paragraph("<b>Skin Type:</b>", bold_style),
            Paragraph(str(report_data.skin_profile.get("skin_type", "N/A")), body_style),
        ],
        [
            Paragraph("<b>Email:</b>", bold_style),
            Paragraph(str(report_data.user_summary.get("email", "N/A")), body_style),
            Paragraph("<b>Primary Concern:</b>", bold_style),
            Paragraph(str(report_data.primary_concern or "None"), body_style),
        ],
        [
            Paragraph("<b>Data Completeness:</b>", bold_style),
            Paragraph(f"{report_data.data_completeness}%", body_style),
            Paragraph("<b>Assessment Confidence:</b>", bold_style),
            Paragraph(f"{report_data.assessment_confidence}%", body_style),
        ]
    ]
    info_table = Table(client_info, colWidths=[1.3 * inch, 2.3 * inch, 1.4 * inch, 2.6 * inch])
    info_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BG_CARD),
        ('BOX', (0, 0), (-1, -1), 0.5, BORDER),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, BORDER),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 8))

    # 3. 5-Factor AuraScore Table
    story.append(Paragraph("1. Overall Dermal Health Score & 5-Factor Breakdown", heading_style))

    score_rows = [
        [
            Paragraph("<b>Factor / Dimension</b>", bold_style),
            Paragraph("<b>Earned Points</b>", bold_style),
            Paragraph("<b>Max Target</b>", bold_style),
            Paragraph("<b>Dermal Health Status</b>", bold_style)
        ]
    ]
    for b in report_data.score_breakdown:
        pct = (b.earned / b.max_possible) if b.max_possible else 1.0
        status_txt = "Optimal" if pct >= 0.8 else ("Moderate" if pct >= 0.5 else "Needs Support")
        score_rows.append([
            Paragraph(b.label, body_style),
            Paragraph(f"<b>{b.earned}</b>", body_style),
            Paragraph(f"{b.max_possible}", body_style),
            Paragraph(status_txt, body_style)
        ])

    score_rows.append([
        Paragraph("<b>COMPOSITE AURASCORE</b>", bold_style),
        Paragraph(f"<b><font color='#0d9488' size=11>{report_data.overall_skin_health_score}</font></b>", bold_style),
        Paragraph("<b>100</b>", bold_style),
        Paragraph("<b>Excellent</b>" if report_data.overall_skin_health_score >= 80 else ("Good" if report_data.overall_skin_health_score >= 60 else "Action Recommended"), bold_style)
    ])

    score_table = Table(score_rows, colWidths=[3.2 * inch, 1.3 * inch, 1.3 * inch, 1.8 * inch])
    score_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f766e")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, BORDER),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#ccfbf1")),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
    ]))
    story.append(score_table)
    story.append(Spacer(1, 8))

    # 4. Personalized Routine Protocol (AM & PM)
    story.append(Paragraph("2. Chronobiology Skincare Regimen Protocol", heading_style))

    routine_rows = [
        [
            Paragraph("<b>Phase</b>", bold_style),
            Paragraph("<b>Step</b>", bold_style),
            Paragraph("<b>Category / Product Type</b>", bold_style),
            Paragraph("<b>Clinical Instructions & Direction</b>", bold_style)
        ]
    ]

    for s in report_data.morning_routine:
        routine_rows.append([
            Paragraph("☀️ Morning", body_style),
            Paragraph(f"Step {s.step_number}", body_style),
            Paragraph(f"<b>{s.product_type}</b>", body_style),
            Paragraph(s.instructions, body_style)
        ])

    for s in report_data.evening_routine:
        routine_rows.append([
            Paragraph("🌙 Evening", body_style),
            Paragraph(f"Step {s.step_number}", body_style),
            Paragraph(f"<b>{s.product_type}</b>", body_style),
            Paragraph(s.instructions, body_style)
        ])

    if report_data.weekly_routine:
        for s in report_data.weekly_routine:
            routine_rows.append([
                Paragraph("✨ Weekly", body_style),
                Paragraph(f"Step {s.step_number}", body_style),
                Paragraph(f"<b>{s.product_type}</b>", body_style),
                Paragraph(s.instructions, body_style)
            ])

    routine_table = Table(routine_rows, colWidths=[1.1 * inch, 0.8 * inch, 2.0 * inch, 3.7 * inch])
    routine_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#334155")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, BORDER),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_CARD]),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(routine_table)
    story.append(Spacer(1, 8))

    # 5. Active Ingredient Recommendations & Safety Notes
    story.append(Paragraph("3. Target Active Ingredients & Safety Exclusions", heading_style))

    ing_text = ", ".join([r.ingredient_or_category for r in report_data.ingredient_recommendations]) if report_data.ingredient_recommendations else "Standard biocompatible active regimen."
    safety_text = "<br/>• ".join(report_data.safety_notes) if report_data.safety_notes else "No active contraindications identified."

    recs_table_data = [
        [
            Paragraph("<b>Target Bioactive Ingredients:</b>", bold_style),
            Paragraph(ing_text, body_style)
        ],
        [
            Paragraph("<b>Safety & Allergy Notes:</b>", bold_style),
            Paragraph(f"• {safety_text}", body_style)
        ]
    ]
    recs_table = Table(recs_table_data, colWidths=[2.2 * inch, 5.4 * inch])
    recs_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BG_CARD),
        ('BOX', (0, 0), (-1, -1), 0.5, BORDER),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, BORDER),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(recs_table)
    story.append(Spacer(1, 8))

    # 6. Actionable Next Steps & Disclaimer
    story.append(Paragraph("4. Recommended Clinical Next Steps", heading_style))
    steps_p = "<br/>".join([f"<b>{i+1}.</b> {step}" for i, step in enumerate(report_data.recommended_next_steps)])
    story.append(Paragraph(steps_p, body_style))
    story.append(Spacer(1, 10))

    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER, spaceBefore=4, spaceAfter=6))
    story.append(Paragraph(
        "<b>Medical Disclaimer:</b> AuraSkin AI Skin Intelligence is a clinical decision-support and personalized cosmetic analysis system. "
        "It does not substitute for licensed medical diagnosis, pathology assessment, or specialized prescription dermatology treatment. "
        "Perform a 24-hour patch test prior to introducing new active cosmeceuticals.",
        disclaimer_style
    ))

    # Build Document
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


# =========================================================================
# EXCEL REPORT EXPORT GENERATION (Openpyxl)
# =========================================================================

def generate_report_excel(db: Session, user: User) -> bytes:
    """
    Generates a structured multi-sheet Excel spreadsheet with assessment, routines, adherence, and progress.
    """
    report_data = generate_skin_intelligence_report(db, user)

    wb = openpyxl.Workbook()

    # Styling helper definitions
    header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    # -------------------------------------------------------------
    # Sheet 1: Executive Summary
    # -------------------------------------------------------------
    ws_exec = wb.active
    ws_exec.title = "AuraSkin Summary"

    ws_exec.append(["AURASKIN AI SKIN INTELLIGENCE — EXECUTIVE REPORT"])
    ws_exec.append(["Generated At", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")])
    ws_exec.append([])

    ws_exec.append(["User Profile"])
    ws_exec.append(["Client Name", report_data.user_summary.get("name")])
    ws_exec.append(["Email", report_data.user_summary.get("email")])
    ws_exec.append(["Skin Type", report_data.skin_profile.get("skin_type")])
    ws_exec.append(["Primary Concern", report_data.primary_concern])
    ws_exec.append(["Secondary Concerns", ", ".join(report_data.secondary_concerns)])
    ws_exec.append(["Composite AuraScore", report_data.overall_skin_health_score])
    ws_exec.append(["Assessment Confidence", f"{report_data.assessment_confidence}%"])
    ws_exec.append(["Data Completeness", f"{report_data.data_completeness}%"])

    # -------------------------------------------------------------
    # Sheet 2: 5-Factor Score Breakdown
    # -------------------------------------------------------------
    ws_score = wb.create_sheet(title="5-Factor Breakdown")
    ws_score.append(["Component Label", "Earned Points", "Max Possible", "Weighting"])
    for cell in ws_score[1]:
        cell.fill = header_fill
        cell.font = header_font

    for b in report_data.score_breakdown:
        ws_score.append([b.label, b.earned, b.max_possible, f"{round((b.max_possible / 100) * 100)}%"])

    ws_score.append(["TOTAL AURASCORE", report_data.overall_skin_health_score, 100, "100%"])

    # -------------------------------------------------------------
    # Sheet 3: Daily Skincare Routines
    # -------------------------------------------------------------
    ws_routine = wb.create_sheet(title="Skincare Regimen")
    ws_routine.append(["Period", "Step #", "Product Category", "Instructions"])
    for cell in ws_routine[1]:
        cell.fill = header_fill
        cell.font = header_font

    for s in report_data.morning_routine:
        ws_routine.append(["Morning", s.step_number, s.product_type, s.instructions])
    for s in report_data.evening_routine:
        ws_routine.append(["Evening", s.step_number, s.product_type, s.instructions])
    if report_data.weekly_routine:
        for s in report_data.weekly_routine:
            ws_routine.append(["Weekly", s.step_number, s.product_type, s.instructions])

    # -------------------------------------------------------------
    # Sheet 4: Routine Adherence History
    # -------------------------------------------------------------
    ws_adherence = wb.create_sheet(title="Adherence Logs")
    ws_adherence.append(["Date (YYYY-MM-DD)", "Morning Completed", "Evening Completed", "Adherence Rate %", "Logged At"])
    for cell in ws_adherence[1]:
        cell.fill = header_fill
        cell.font = header_font

    adherences = db.query(RoutineAdherenceRecord).filter(
        RoutineAdherenceRecord.user_id == user.id
    ).order_by(RoutineAdherenceRecord.date.desc()).all()

    for a in adherences:
        ws_adherence.append([
            a.date,
            "YES" if a.morning_completed else "NO",
            "YES" if a.evening_completed else "NO",
            f"{round(a.adherence_rate * 100, 1)}%",
            a.created_at.strftime("%Y-%m-%d %H:%M") if a.created_at else ""
        ])

    # -------------------------------------------------------------
    # Sheet 5: Progress Snapshots
    # -------------------------------------------------------------
    ws_prog = wb.create_sheet(title="Progress Snapshots")
    ws_prog.append(["Snapshot ID", "Total Score", "Skin Condition", "Lifestyle", "Sleep", "Routine", "Hydration", "Barrier Status", "Recorded At"])
    for cell in ws_prog[1]:
        cell.fill = header_fill
        cell.font = header_font

    progress_records = db.query(ProgressRecord).filter(
        ProgressRecord.user_id == user.id
    ).order_by(ProgressRecord.recorded_at.desc()).all()

    for p in progress_records:
        ws_prog.append([
            p.id,
            p.total_score,
            p.skin_condition_score,
            p.lifestyle_score,
            p.sleep_score,
            p.routine_consistency_score,
            p.hydration_score,
            p.barrier_status,
            p.recorded_at.strftime("%Y-%m-%d %H:%M") if p.recorded_at else ""
        ])

    # Format column widths automatically across all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
